from pathlib import Path

import pytest
from docx import Document

from app.services.doc_generator import BidDocumentGenerator


def _generator(tmp_path):
    templates_dir = tmp_path / "templates"
    output_dir = tmp_path / "output"
    templates_dir.mkdir()
    return BidDocumentGenerator(templates_dir=templates_dir, output_dir=output_dir)


def test_output_dir_created(tmp_path):
    generator = _generator(tmp_path)
    assert generator.output_dir.is_dir()


def test_partner_count_rejects_missing_middle_partner(tmp_path):
    generator = _generator(tmp_path)
    with pytest.raises(ValueError, match="first partner"):
        generator.determine_partner_count(
            {
                "LEAD_PARTNER_NAME": "Lead Co",
                "FIRST_PARTNER_NAME": "",
                "SECOND_PARTNER_NAME": "Second Co",
            }
        )


def test_partner_count_supports_one_two_and_three_partners(tmp_path):
    generator = _generator(tmp_path)
    assert generator.determine_partner_count({"LEAD_PARTNER_NAME": "Lead"}) == 1
    assert generator.determine_partner_count({"LEAD_PARTNER_NAME": "Lead", "FIRST_PARTNER_NAME": "First"}) == 2
    assert (
        generator.determine_partner_count(
            {"LEAD_PARTNER_NAME": "Lead", "FIRST_PARTNER_NAME": "First", "SECOND_PARTNER_NAME": "Second"}
        )
        == 3
    )


def test_unresolved_placeholder_detection(tmp_path):
    generator = _generator(tmp_path)
    document = Document()
    document.add_paragraph("Hello {{KNOWN}} and {{MISSING}}")
    generator.replace_in_document(document, {"{{KNOWN}}": "World"})
    assert generator.unresolved_placeholders(document) == ["{{MISSING}}"]


def test_select_template_by_partner_count(tmp_path):
    generator = _generator(tmp_path)
    (generator.templates_dir / "master_template_1.docx").touch()
    (generator.templates_dir / "master_template_2.docx").touch()
    (generator.templates_dir / "master_template_3.docx").touch()
    assert generator.select_template(1) == "master_template_1.docx"
    assert generator.select_template(2) == "master_template_2.docx"
    assert generator.select_template(3) == "master_template_3.docx"


def test_select_template_falls_back_when_template_1_missing(tmp_path):
    generator = _generator(tmp_path)
    (generator.templates_dir / "master_template_2.docx").touch()
    assert generator.select_template(1) == "master_template_2.docx"


REAL_TEMPLATES = Path(__file__).resolve().parents[1] / "templates"
SIG_ALT_RE = r"(AUTHORISED_SIG|_CEO_SIG|_STAMP|_PARTNER_MD\d|_MD\d_SIG)"


def _image_slot_alts(docx_path):
    """Alt texts of every image left in the generated document (body + headers/footers)."""
    import re
    import zipfile

    alts = []
    with zipfile.ZipFile(docx_path) as z:
        for name in z.namelist():
            if re.match(r"word/(document|header\d*|footer\d*)\.xml$", name):
                xml = z.read(name).decode("utf-8", "ignore")
                alts += re.findall(r'<wp:docPr [^>]*descr="([^"]*)"', xml)
    return alts


def _bid_data(partners: int) -> dict:
    data = {
        "BID_TYPE": "Joint Venture",
        "JV_NAME": "ABC - XYZ J/V",
        "PROJECT_NAME": "Road Project",
        "EMPLOYER_NAME": "Employer",
        "EMPLOYER_ADDRESS": "Kathmandu",
        "JV_ADDRESS": "Dharan",
        "IFB_NUMBER": "NCB/W/01",
        "BID_DATE": "2026-09-28",
        "BID_VALIDITY_PERIOD": "120 days",
        "LEAD_PARTNER_NAME": "Lead Co",
        "LEAD_PARTNER_SHORT": "LC",
        "LEAD_ADDRESS": "Lead Address",
        "LEAD_PARTNER_CEO": "Lead CEO",
        "AUTHORIZED_PERSON_NAME": "Lead CEO",
        "L_PER": "100",
    }
    if partners >= 2:
        data.update({"FIRST_PARTNER_NAME": "First Co", "FIRST_PARTNER_SHORT": "FC", "FIRST_ADDRESS": "First Address", "FIRST_PARTNER_CEO": "First CEO", "L_PER": "60", "F_PER": "40"})
    if partners == 3:
        data.update({"SECOND_PARTNER_NAME": "Second Co", "SECOND_PARTNER_SHORT": "SC", "SECOND_ADDRESS": "Second Address", "SECOND_PARTNER_CEO": "Second CEO", "F_PER": "20", "S_PER": "20"})
    return data


def _tiny_png() -> bytes:
    import struct
    import zlib

    def chunk(kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    pixels = zlib.compress(bytes(5))
    return bytes.fromhex("89504e470d0a1a0a") + chunk(b"IHDR", header) + chunk(b"IDAT", pixels) + chunk(b"IEND", b"")


@pytest.mark.parametrize("partners", [1, 2, 3])
def test_generates_without_any_signature_or_stamp(tmp_path, partners):
    """Signatures are optional: with no uploads the doc still generates and no sample images remain."""
    import re

    from app.services.bid_compute import with_derived_fields
    from app.services.validation import validate_bid

    data = with_derived_fields(_bid_data(partners))
    assert validate_bid(data, authorized_signature_present=False) == []

    generator = BidDocumentGenerator(templates_dir=REAL_TEMPLATES, output_dir=tmp_path / "out")
    output = generator.generate(data, image_mapping={})

    leftover = [alt for alt in _image_slot_alts(output) if re.search(SIG_ALT_RE, alt)]
    assert leftover == []


def test_uploaded_signature_is_kept(tmp_path):
    """Slots that do have an upload keep their image (only missing ones are removed)."""
    import re

    from app.services.bid_compute import with_derived_fields

    sig = tmp_path / "sig.png"
    sig.write_bytes(_tiny_png())
    data = with_derived_fields(_bid_data(2))
    mapping = {"LEAD_CEO_SIG": str(sig), "AUTHORISED_SIG": str(sig)}

    generator = BidDocumentGenerator(templates_dir=REAL_TEMPLATES, output_dir=tmp_path / "out")
    alts = _image_slot_alts(generator.generate(data, image_mapping=mapping))

    kept = {alt for alt in alts if re.search(SIG_ALT_RE, alt)}
    assert kept == {"LEAD_CEO_SIG", "AUTHORISED_SIG"}
